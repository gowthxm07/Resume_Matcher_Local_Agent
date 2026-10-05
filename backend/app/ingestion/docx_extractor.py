"""
DOCX extraction module utilizing python-docx.
Extracts paragraphs, table contents, and document properties safely.
"""

import io
import time
from pathlib import Path
from typing import Union
import docx
from app.ingestion.document import NormalizedDocument
from app.core.logging import logger


class DOCXExtractor:
    """Extracts text content and metadata from Word .docx files."""

    @classmethod
    def extract(cls, file_path_or_bytes: Union[str, Path, bytes], filename: str) -> NormalizedDocument:
        """
        Extract text and metadata from a DOCX document.

        Args:
            file_path_or_bytes: Path to DOCX file or raw bytes
            filename: Original filename

        Returns:
            NormalizedDocument instance
        """
        start_time = time.perf_counter()
        try:
            if isinstance(file_path_or_bytes, (str, Path)):
                path = Path(file_path_or_bytes)
                if not path.exists():
                    raise FileNotFoundError(f"DOCX file not found at: {path}")
                doc = docx.Document(str(path))
                file_size = path.stat().st_size
            else:
                doc = docx.Document(io.BytesIO(file_path_or_bytes))
                file_size = len(file_path_or_bytes)

            content_lines = []

            # 1. Extract paragraphs
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    content_lines.append(text)

            # 2. Extract tables (resumes often use tables for layouts)
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    # Deduplicate adjacent cells caused by merged table cells
                    deduped_cells = []
                    for c in row_cells:
                        if not deduped_cells or deduped_cells[-1] != c:
                            deduped_cells.append(c)
                    if deduped_cells:
                        content_lines.append(" | ".join(deduped_cells))

            extracted_text = "\n\n".join(content_lines).strip()

            # Metadata from core properties
            doc_metadata = {}
            try:
                cp = doc.core_properties
                doc_metadata = {
                    "file_size_bytes": file_size,
                    "title": cp.title or None,
                    "author": cp.author or None,
                    "created": str(cp.created) if cp.created else None,
                    "modified": str(cp.modified) if cp.modified else None,
                    "table_count": len(doc.tables),
                    "paragraph_count": len(doc.paragraphs),
                }
            except Exception as prop_err:
                logger.debug(f"Could not read DOCX core properties: {prop_err}")
                doc_metadata = {"file_size_bytes": file_size}

            elapsed_ms = (time.perf_counter() - start_time) * 1000.0
            logger.info(
                f"Successfully extracted DOCX '{filename}' ({len(extracted_text)} chars) in {elapsed_ms:.1f}ms"
            )

            return NormalizedDocument(
                filename=filename,
                document_type="docx",
                extracted_text=extracted_text,
                page_count=None,  # DOCX doesn't store fixed pagination without rendering
                extraction_time_ms=round(elapsed_ms, 2),
                metadata=doc_metadata,
            )

        except Exception as exc:
            logger.error(f"Error extracting DOCX '{filename}': {exc}")
            raise
